import base64
import io
import json
import mimetypes
import os
from pathlib import Path
from typing import List, Dict, Any

import boto3
import docx
from PIL import Image
import fitz  # PyMuPDF
from botocore.exceptions import ClientError
from langchain.text_splitter import RecursiveCharacterTextSplitter
from loguru import logger
from dotenv import load_dotenv

# Constants
IMAGE_DIR = Path("rag/knowledge/images")
IMAGE_DIR.mkdir(parents=True, exist_ok=True)

class DocumentProcessor:
    """
    Handles the processing of documents, including text and image extraction,
    embedding generation, and preparing data for the knowledge base.
    """

    def __init__(self, mock: bool = False):
        """
        Initializes the DocumentProcessor.

        Args:
            mock (bool): If True, runs in mock mode without actual API calls.
        """
        # Load environment variables from .env file
        load_dotenv()
        
        self.mock = mock
        if not self.mock:
            self.bedrock_client = boto3.client("bedrock-runtime", region_name=os.environ.get("AWS_REGION", "ap-southeast-2"))
            self.s3_client = boto3.client("s3")
            self.s3_bucket_name = os.environ.get("AWS_S3_BUCKET_NAME","reco-demo-res")
            if not self.s3_bucket_name:
                raise ValueError("AWS_S3_BUCKET_NAME environment variable not set.")
        else:
            self.bedrock_client = None
            self.s3_client = None
            self.s3_bucket_name = None
            logger.info("DocumentProcessor running in mock mode.")
            
        self.embedding_model_id = "amazon.titan-embed-text-v2:0"
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=200,
            length_function=len
        )

    def process_document(self, file_path: str) -> List[Dict[str, Any]]:
        """
        Main function to process a single document.
        """
        logger.info(f"Processing document: {file_path}")
        raw_text, extracted_images = self._extract_text_and_images(file_path)
        processed_text = self._describe_images_and_insert_placeholders(raw_text, extracted_images)
        logger.info(f"Processed text: {processed_text}")
        text_chunks = self.text_splitter.split_text(processed_text)
        processed_chunks = self._generate_embeddings(text_chunks)
        logger.info(f"Successfully processed {len(processed_chunks)} chunks from {file_path}")
        return processed_chunks

    def _upload_image_to_s3(self, image_bytes: bytes, image_filename: str) -> str:
        """Uploads image bytes to S3 and returns the public URL."""
        if self.mock:
            logger.info(f"Mock uploading {image_filename} to S3.")
            # In mock mode, we don't upload but can return a predictable fake URL
            return f"https://fake-bucket.s3.amazonaws.com/images/{image_filename}"

        content_type = mimetypes.guess_type(image_filename)[0] or 'application/octet-stream'
        s3_object_name = f"images/{image_filename}"
        
        try:
            self.s3_client.head_object(Bucket=self.s3_bucket_name, Key=s3_object_name)
            logger.info(f"Image already exists in S3: s3://{self.s3_bucket_name}/{s3_object_name}")
        except ClientError as e:
            if e.response['Error']['Code'] == '404':
                self.s3_client.upload_fileobj(
                    io.BytesIO(image_bytes),
                    self.s3_bucket_name,
                    s3_object_name,
                    ExtraArgs={'ACL': 'public-read', 'ContentType': content_type}
                )
                logger.info(f"Uploaded image to S3: s3://{self.s3_bucket_name}/{s3_object_name}")
            else:
                logger.error(f"Failed to check/upload image to S3: {e}", exc_info=True)
                raise
        
        # Return the public URL of the image
        return f"https://{self.s3_bucket_name}.s3.amazonaws.com/{s3_object_name}"

    def _extract_text_and_images(self, file_path: str) -> (str, List[Dict[str, Any]]):
        """
        Extracts text and images from a given document (PDF or DOCX).
        Images are handled in memory and passed to subsequent functions.
        """
        file_extension = Path(file_path).suffix.lower()
        if file_extension == ".pdf":
            return self._extract_from_pdf(file_path)
        elif file_extension == ".docx":
            return self._extract_from_docx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_extension}")

    def _extract_from_pdf(self, file_path: str) -> (str, List[Dict[str, Any]]):
        """Extracts text and image bytes from a PDF file."""
        logger.info(f"Extracting from PDF: {file_path}")
        text = ""
        extracted_images = []
        doc = fitz.open(file_path)

        for page_num, page in enumerate(doc):
            text += page.get_text()
            image_list = page.get_images(full=True)
            for img_index, img in enumerate(image_list):
                xref = img[0]
                base_image = doc.extract_image(xref)
                image_bytes = base_image["image"]
                image_ext = base_image["ext"]
                
                image_filename = f"pdf_{Path(file_path).stem}_p{page_num+1}_img{img_index+1}.{image_ext}"
                placeholder = f"\n[image_placeholder:{image_filename}]\n"
                text += placeholder
                
                extracted_images.append({"bytes": image_bytes, "filename": image_filename})

        doc.close()
        logger.info(f"Extracted {len(extracted_images)} images from {file_path}")
        return text, extracted_images

    def _extract_from_docx(self, file_path: str) -> (str, List[Dict[str, Any]]):
        """Extracts text and image bytes from a DOCX file, preserving their order."""
        logger.info(f"Extracting from DOCX: {file_path}")
        doc = docx.Document(file_path)
        text_content = ""
        extracted_images = []
        
        image_part_map = {
            rId: rel.target_part
            for rId, rel in doc.part.rels.items()
            if "image" in rel.target_ref
        }
        
        img_counter = 0
        NAMESPACES = {
            'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
            'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
            'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
        }

        for para in doc.paragraphs:
            para_content = ""
            for run in para.runs:
                drawing_elems = run.element.findall('.//w:drawing', namespaces=NAMESPACES)
                if drawing_elems:
                    for drawing in drawing_elems:
                        blip_elems = drawing.findall('.//a:blip', namespaces=NAMESPACES)
                        for blip in blip_elems:
                            rId = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                            if rId and rId in image_part_map:
                                image_part = image_part_map[rId]
                                image_bytes = image_part.blob
                                
                                img_counter += 1
                                image_ext = image_part.content_type.split('/')[-1]
                                if image_ext == "jpeg": image_ext = "jpg"
                                
                                image_filename = f"docx_{Path(file_path).stem}_img{img_counter}.{image_ext}"
                                placeholder = f"\n[image_placeholder:{image_filename}]\n"
                                para_content += placeholder
                                extracted_images.append({"bytes": image_bytes, "filename": image_filename})
                else:
                    para_content += run.text
            text_content += para_content + "\n"

        logger.info(f"Extracted {len(extracted_images)} images from {file_path}")
        return text_content, extracted_images

    def _describe_images_and_insert_placeholders(self, text: str, images: List[Dict[str, Any]]) -> str:
        """
        Uploads images to S3, gets descriptions, and inserts info into text.
        """
        if not images:
            return text

        logger.info(f"Describing and processing {len(images)} images...")
        for image_data in images:
            image_bytes = image_data["bytes"]
            image_filename = image_data["filename"]

            # Upload to S3 and get public URL
            s3_url = self._upload_image_to_s3(image_bytes, image_filename)
            
            # Get image description from bytes
            description = self._get_image_description(image_bytes)
            
            image_info = {
                "description": description,
                "imgpath": s3_url
            }
            info_str = f"[image_info]{json.dumps(image_info)}[/image_info]"
            
            placeholder = f"[image_placeholder:{image_filename}]"
            text = text.replace(placeholder, info_str)
            
        return text

    def _get_image_description(self, image_bytes: bytes) -> str:
        """
        Generates a text description for a single image's bytes using a multimodal LLM.
        """
        if self.mock:
            logger.info(f"Getting mock description for an image...")
            return "A mock description for the provided image."
        
        logger.info(f"Getting image description (real)...")
        try:
            base64_image = base64.b64encode(image_bytes).decode("utf-8")
            
            # Infer image format from bytes if possible, otherwise default
            try:
                img = Image.open(io.BytesIO(image_bytes))
                image_format = img.format.lower() if img.format else 'png'
            except Exception:
                image_format = 'png' # Default if format detection fails

            content = [
                {
                    "image": {
                        "format": image_format,
                        "source": { "bytes": base64_image }
                    }
                },
                {"text": "Describe this image in detail."}
            ]
            request_body = {
                "messages": [{"role": "user", "content": content}],
                "inferenceConfig": {"max_new_tokens": 300, "temperature": 0.5, "top_p": 0.9}
            }
            response_body = self._invoke_bedrock("apac.amazon.nova-lite-v1:0", request_body)
            return response_body.get('output', {}).get('message', {}).get('content', [{}])[0].get('text', '')

        except Exception as e:
            logger.exception(f"Error getting image description: {e}")
            return f"Error describing image: {e}"

    def _generate_embeddings(self, text_chunks: List[str]) -> List[Dict[str, Any]]:
        """
        Generates vector embeddings for a list of text chunks using Bedrock.
        """
        if self.mock:
            logger.info(f"Generating embeddings for {len(text_chunks)} chunks (mock)...")
            return [{"text": chunk, "embedding": [0.0] * 1536, "metadata": {}} for chunk in text_chunks]

        logger.info(f"Generating embeddings for {len(text_chunks)} chunks (real)...")
        processed_chunks = []
        for chunk in text_chunks:
            try:
                body = json.dumps({"inputText": chunk})
                response = self.bedrock_client.invoke_model(
                    body=body,
                    modelId=self.embedding_model_id,
                    accept="application/json",
                    contentType="application/json"
                )
                response_body = json.loads(response.get("body").read())
                embedding = response_body.get("embedding")
                processed_chunks.append({"text": chunk, "embedding": embedding, "metadata": {}})
            except Exception as e:
                logger.exception(f"Error generating embedding for chunk: {e}")
                continue
        return processed_chunks
    
    def _invoke_bedrock(self, model_id: str, body: Dict[str, Any]) -> Dict[str, Any]:
        """Generic method to invoke a Bedrock model."""
        try:
            response = self.bedrock_client.invoke_model(
                body=json.dumps(body),
                modelId=model_id,
                contentType='application/json',
                accept='application/json'
            )
            return json.loads(response.get('body').read())
        except Exception as e:
            logger.error(f"A Bedrock client error occurred when invoking {model_id}: {e}")
            raise

if __name__ == '__main__':
    logger.info("Running document processor example...")
    # Set env vars for local testing if not present
    os.environ.setdefault("AWS_REGION", "ap-southeast-2")
    os.environ.setdefault("AWS_S3_BUCKET_NAME", "your-test-bucket")
    
    processor = DocumentProcessor(mock=True) # Using mock to avoid real API calls in example
    dummy_file = "dummy.txt"
    with open(dummy_file, "w") as f:
        f.write("This is a test document.")
    
    try:
        processed_data = processor.process_document(dummy_file)
        logger.info(f"Processed data: {processed_data}")
    except ValueError as e:
        logger.warning(f"Caught expected error for unsupported file type: {e}")

    os.remove(dummy_file)
    logger.info("Example run finished.")